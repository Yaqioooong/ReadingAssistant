"""表格结构化测试：docx 表格 → Markdown 原子块，分块不腰斩。"""

from pathlib import Path

from docx import Document

from reading_assistant.parsers.base import TABLE_BEGIN, TABLE_END, Chapter, ParsedBook
from reading_assistant.parsers.docx_parser import DocxParser
from reading_assistant.rag.chunking import _MAX_TABLE_CHARS, chunk_book


class TestChunkTableBlocks:
    def test_table_is_atomic_chunk(self) -> None:
        content = (
            '第一章开头段落。\n\n'
            f'{TABLE_BEGIN}\n'
            '| 姓名 | 职务 | 年份 |\n'
            '| --- | --- | --- |\n'
            '| 少白公 | 国家元首 | 1966 |\n'
            f'{TABLE_END}\n\n'
            '表格之后的正文。'
        )
        book = ParsedBook(title='T', chapters=[Chapter('第一章', content)])
        chunks = chunk_book(book, chunk_size=50, chunk_overlap=0)
        assert any(c.metadata.get('block_type') == 'table' for c in chunks)

    def test_table_rows_not_split(self) -> None:
        # 即使 chunk_size 很小，表格也必须整块保留（不被从行中间切断）
        content = (
            f'{TABLE_BEGIN}\n'
            '| 列一 | 列二 | 列三 |\n'
            '| --- | --- | --- |\n'
            '| A1 | B1 | C1 |\n'
            '| A2 | B2 | C2 |\n'
            f'{TABLE_END}\n'
        )
        book = ParsedBook(title='T', chapters=[Chapter('正文', content)])
        chunks = chunk_book(book, chunk_size=20, chunk_overlap=0)
        tables = [c for c in chunks if c.metadata.get('block_type') == 'table']
        assert len(tables) == 1
        text = tables[0].text
        assert '| A1 | B1 | C1 |' in text
        assert '| A2 | B2 | C2 |' in text
        assert '| --- | --- | --- |' in text
        assert TABLE_BEGIN not in text and TABLE_END not in text  # 哨兵不进正文

    def test_text_segments_around_table_unchanged(self) -> None:
        content = (
            '前文' + '字' * 100 + '\n\n'
            f'{TABLE_BEGIN}\n| A |\n| --- |\n| 1 |\n{TABLE_END}\n\n'
            '后文' + '字' * 100
        )
        book = ParsedBook(title='T', chapters=[Chapter('第一章', content)])
        chunks = chunk_book(book, chunk_size=60, chunk_overlap=0)
        pre_idx = next(i for i, c in enumerate(chunks) if c.text.startswith('前文'))
        table_idx = next(
            i for i, c in enumerate(chunks) if c.metadata.get('block_type') == 'table'
        )
        post_idx = next(i for i, c in enumerate(chunks) if '后文' in c.text)
        assert pre_idx < table_idx < post_idx  # 顺序：前文 → 表格 → 后文
        assert all(not c.text.startswith(TABLE_BEGIN) for c in chunks)


class TestDocxTable:
    def test_docx_table_to_atomic_chunk(self, tmp_path: Path) -> None:
        doc_path = tmp_path / 'table.docx'
        doc = Document()
        doc.add_heading('第一章 数据', level=1)
        doc.add_paragraph('表格前的说明文字。')
        table = doc.add_table(rows=3, cols=3)
        data = [['姓名', '职务', '年份'], ['少白公', '国家元首', '1966'], ['舵手', '领袖', '1969']]
        for r, row in enumerate(data):
            for c, value in enumerate(row):
                table.cell(r, c).text = value
        doc.add_paragraph('表格后的总结文字。')
        doc.save(str(doc_path))

        parsed = DocxParser().parse(str(doc_path))
        content = parsed.chapters[0].content
        assert '| 姓名 | 职务 | 年份 |' in content
        assert TABLE_BEGIN in content and TABLE_END in content

        chunks = chunk_book(parsed)
        tables = [c for c in chunks if c.metadata.get('block_type') == 'table']
        assert len(tables) == 1
        assert '少白公' in tables[0].text and '1966' in tables[0].text
        assert '国家元首' in tables[0].text


def _make_table(rows: int, cols: int = 6) -> str:
    """构造 Markdown 表格（表头 + 分隔行 + rows 行数据）。"""
    header = '| ' + ' | '.join(f'列{c}' for c in range(cols)) + ' |'
    separator = '| ' + ' | '.join(['---'] * cols) + ' |'
    body = [
        '| ' + ' | '.join(f'数据{r}-{c}' for c in range(cols)) + ' |'
        for r in range(rows)
    ]
    return '\n'.join([header, separator, *body])


def _book_with_table(table: str) -> ParsedBook:
    content = f'表格前的正文。\n\n{TABLE_BEGIN}\n{table}\n{TABLE_END}\n\n表格后的正文。'
    return ParsedBook(title='T', chapters=[Chapter('第一章', content)])


def _table_chunks(table: str):
    return [
        c
        for c in chunk_book(_book_with_table(table))
        if c.metadata.get('block_type') == 'table'
    ]


class TestTableSizeLimit:
    """P0-2：表格原子块必须有上限。

    实测故障：5000 行表 → 499,013 字符单 chunk = **623.8x chunk_size**；
    超过 text-embedding-v4 的 8192 token 上限（20000 字符 → 400 InvalidParameter）
    → ``embed_documents`` 批量调用报错 → **整本书入库失败**。

    修法遵循业界做法（LlamaIndex table node parser / unstructured）：可以切，
    但**每一批都必须重复表头 + 分隔行**，否则第二块起彻底丢失列名语义。
    """

    def test_small_table_stays_atomic(self) -> None:
        table = _make_table(rows=3, cols=3)
        assert len(table) <= _MAX_TABLE_CHARS

        tables = _table_chunks(table)

        assert len(tables) == 1, '未超限的表格应保持原子（不切）'
        assert '列0' in tables[0].text

    def test_large_table_is_split(self) -> None:
        table = _make_table(rows=400, cols=6)
        assert len(table) > _MAX_TABLE_CHARS

        tables = _table_chunks(table)

        assert len(tables) > 1, '超限表格必须被切分'

    def test_every_part_carries_header_and_separator(self) -> None:
        """核心原则：表头必须每块都带 —— 否则列名语义丢失。"""
        tables = _table_chunks(_make_table(rows=400, cols=6))

        assert len(tables) > 1
        for part in tables:
            lines = part.text.splitlines()
            assert lines[0].startswith('| 列0 |'), f'批次缺表头: {lines[0][:48]}'
            assert lines[1].startswith('| --- |'), '批次缺分隔行'

    def test_no_part_exceeds_limit(self) -> None:
        tables = _table_chunks(_make_table(rows=400, cols=6))

        for part in tables:
            assert len(part.text) <= _MAX_TABLE_CHARS, f'批次超限: {len(part.text)}'

    def test_rows_are_never_split_mid_row(self) -> None:
        tables = _table_chunks(_make_table(rows=400, cols=6))

        for part in tables:
            for line in part.text.splitlines():
                assert line.startswith('|') and line.endswith('|'), f'行被腰斩: {line[:48]}'

    def test_all_parts_keep_block_type(self) -> None:
        tables = _table_chunks(_make_table(rows=400, cols=6))

        assert len(tables) >= 2, '每个批次都必须仍是 table 类型'

    def test_no_row_is_lost(self) -> None:
        """切分不得丢数据：所有原行都应出现在某个批次里。"""
        tables = _table_chunks(_make_table(rows=400, cols=6))
        joined = '\n'.join(part.text for part in tables)

        for row in range(400):
            assert f'数据{row}-0' in joined, f'第 {row} 行在切分中丢失'

    def test_regression_5000_row_table_no_longer_unbounded(self) -> None:
        """原始故障规模：5000 行表 → 623.8x chunk_size。

        修复后无任何单块超过 _MAX_TABLE_CHARS，嵌入输入回到 8192 token 以内。
        """
        table = _make_table(rows=5000, cols=6)

        tables = _table_chunks(table)

        assert len(tables) > 1
        largest = max(len(part.text) for part in tables)
        assert largest <= _MAX_TABLE_CHARS, f'仍有超限单块: {largest}'
        assert largest < len(table), '切分后单块必须显著小于整表'

    def test_single_huge_row_falls_back_to_cell_split(self) -> None:
        """单行本身超限 → 单元格边界兜底，且每片仍带表头。"""
        cols = 200
        header = '| ' + ' | '.join(f'列{c}' for c in range(cols)) + ' |'
        separator = '| ' + ' | '.join(['---'] * cols) + ' |'
        row = '| ' + ' | '.join('长内容' * 20 for _ in range(cols)) + ' |'
        table = '\n'.join([header, separator, row])
        assert len(table) > _MAX_TABLE_CHARS

        tables = _table_chunks(table)

        assert len(tables) > 1, '超宽行必须被切分，否则仍会撞嵌入上限'
        for part in tables:
            assert part.text.splitlines()[0].startswith('| 列0 |'), '兜底分片也必须带表头'
