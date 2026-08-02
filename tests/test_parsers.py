"""P1 解析器测试：txt / epub / pdf / docx 的 happy path 与边界用例。"""

from io import BytesIO
from pathlib import Path

import pytest
from docx import Document
from ebooklib import epub
from pypdf import PdfReader, PdfWriter

from reading_assistant.parsers import (
    DocxParser,
    EncryptedFileError,
    EpubParser,
    ParseError,
    PdfParser,
    TxtParser,
    UnsupportedFormatError,
    get_parser,
    parse_book,
)


def _make_pdf(text: str = 'Hello PDF') -> bytes:
    """手工构造一个包含一行文本的最小合法 PDF。"""
    content = f'BT /F1 12 Tf 72 720 Td ({text}) Tj ET'.encode('latin-1')
    objects = [
        b'<< /Type /Catalog /Pages 2 0 R >>',
        b'<< /Type /Pages /Kids [3 0 R] /Count 1 >>',
        b'<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] '
        b'/Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>',
        b'<< /Length %d >>\nstream\n%s\nendstream' % (len(content), content),
        b'<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>',
    ]
    out = bytearray(b'%PDF-1.4\n')
    offsets = [0]
    for index, obj in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f'{index} 0 obj\n'.encode() + obj + b'\nendobj\n'
    xref_pos = len(out)
    out += f'xref\n0 {len(objects) + 1}\n'.encode()
    out += b'0000000000 65535 f \n'
    for offset in offsets[1:]:
        out += f'{offset:010d} 00000 n \n'.encode()
    out += (
        f'trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref_pos}\n%%EOF\n'
    ).encode()
    return bytes(out)


def _make_encrypted_pdf(path: Path) -> None:
    reader = PdfReader(BytesIO(_make_pdf()))
    writer = PdfWriter()
    writer.append(reader)
    writer.encrypt('secret')
    with open(path, 'wb') as f:
        writer.write(f)


def _make_epub(path: Path) -> None:
    book = epub.EpubBook()
    book.set_identifier('test-0001')
    book.set_title('测试图书')
    book.set_language('zh')
    book.add_author('作者甲')
    c1 = epub.EpubHtml(title='第一章 开始', file_name='chap_01.xhtml', lang='zh')
    c1.content = '<h1>第一章 开始</h1><p>第一章的内容。</p>'
    c2 = epub.EpubHtml(title='第二章 发展', file_name='chap_02.xhtml', lang='zh')
    c2.content = '<h1>第二章 发展</h1><p>第二章的内容。</p>'
    book.add_item(c1)
    book.add_item(c2)
    book.toc = (c1, c2)
    book.add_item(epub.EpubNcx())
    book.add_item(epub.EpubNav())
    book.spine = ['nav', c1, c2]
    epub.write_epub(str(path), book)


def _make_docx(path: Path, with_author: bool = True) -> None:
    doc = Document()
    doc.core_properties.title = '测试文档'
    if with_author:
        doc.core_properties.author = '作者乙'
    doc.add_heading('第一章 绪论', level=1)
    doc.add_paragraph('绪论正文内容。')
    doc.add_heading('第二章 方法', level=1)
    doc.add_paragraph('方法正文内容。')
    doc.save(str(path))


class TestTxtParser:
    def test_parse_with_chapters(self, tmp_path: Path) -> None:
        book = tmp_path / '三体.txt'
        book.write_text('第一章 开端\n正文一。\n第二章 发展\n正文二。\n', encoding='utf-8')

        result = parse_book(book)

        assert isinstance(get_parser(book), TxtParser)
        assert result.title == '三体'
        assert result.chapter_count == 2
        assert result.chapters[0].title == '第一章 开端'
        assert '正文一。' in result.chapters[0].content
        assert result.chapters[1].title == '第二章 发展'
        assert '正文二。' in result.full_text

    def test_parse_without_chapter_headings(self, tmp_path: Path) -> None:
        book = tmp_path / 'notes.txt'
        book.write_text('只是一段没有章节标题的正文。', encoding='utf-8')

        result = TxtParser().parse(book)

        assert result.chapter_count == 1
        assert result.chapters[0].title == '正文'

    def test_fallback_encoding_for_non_utf8_file(self, tmp_path: Path) -> None:
        book = tmp_path / 'legacy.txt'
        book.write_bytes(b'\xff\xfe\x00\x80legacy text')

        result = TxtParser().parse(book)

        assert result.title == 'legacy'
        assert result.full_text

    def test_empty_file_raises(self, tmp_path: Path) -> None:
        book = tmp_path / 'empty.txt'
        book.write_text('', encoding='utf-8')

        with pytest.raises(ParseError, match='空'):
            TxtParser().parse(book)


class TestEpubParser:
    def test_parse_with_metadata_and_chapters(self, tmp_path: Path) -> None:
        book = tmp_path / 'sample.epub'
        _make_epub(book)

        result = parse_book(book)

        assert isinstance(get_parser(book), EpubParser)
        assert result.title == '测试图书'
        assert result.author == '作者甲'
        assert result.chapter_count == 2
        assert result.chapters[0].title == '第一章 开始'
        assert '第一章的内容。' in result.chapters[0].content

    def test_corrupted_file_raises(self, tmp_path: Path) -> None:
        book = tmp_path / 'broken.epub'
        book.write_bytes(b'not an epub file')

        with pytest.raises(ParseError, match='EPUB'):
            EpubParser().parse(book)


class TestPdfParser:
    def test_parse_extracts_text(self, tmp_path: Path) -> None:
        book = tmp_path / 'sample.pdf'
        book.write_bytes(_make_pdf())

        result = parse_book(book)

        assert isinstance(get_parser(book), PdfParser)
        assert result.title == 'sample'
        assert 'Hello PDF' in result.full_text

    def test_encrypted_file_raises(self, tmp_path: Path) -> None:
        book = tmp_path / 'locked.pdf'
        _make_encrypted_pdf(book)

        with pytest.raises(EncryptedFileError):
            PdfParser().parse(book)

    def test_corrupted_file_raises(self, tmp_path: Path) -> None:
        book = tmp_path / 'broken.pdf'
        book.write_bytes(b'not a pdf')

        with pytest.raises(ParseError, match='PDF'):
            PdfParser().parse(book)


class TestDocxParser:
    def test_parse_with_metadata_and_heading_chapters(self, tmp_path: Path) -> None:
        book = tmp_path / 'sample.docx'
        _make_docx(book)

        result = parse_book(book)

        assert isinstance(get_parser(book), DocxParser)
        assert result.title == '测试文档'
        assert result.author == '作者乙'
        assert result.chapter_count == 2
        assert result.chapters[0].title == '第一章 绪论'
        assert '绪论正文内容。' in result.chapters[0].content

    def test_corrupted_file_raises(self, tmp_path: Path) -> None:
        book = tmp_path / 'broken.docx'
        book.write_bytes(b'not a docx')

        with pytest.raises(ParseError, match='DOCX'):
            DocxParser().parse(book)

    def test_empty_document_raises(self, tmp_path: Path) -> None:
        book = tmp_path / 'empty.docx'
        Document().save(str(book))

        with pytest.raises(ParseError, match='无'):
            DocxParser().parse(book)


class TestFactory:
    def test_extension_dispatch_and_case_insensitive(self, tmp_path: Path) -> None:
        assert isinstance(get_parser('book.TXT'), TxtParser)
        assert isinstance(get_parser('book.EPUB'), EpubParser)
        assert isinstance(get_parser('book.PDF'), PdfParser)
        assert isinstance(get_parser('book.DOCX'), DocxParser)

    def test_unsupported_extension_raises(self) -> None:
        with pytest.raises(UnsupportedFormatError):
            get_parser('book.mobi')


class TestPathResolution:
    def test_missing_file_raises_clear_error(self, tmp_path: Path, monkeypatch) -> None:
        monkeypatch.chdir(tmp_path)

        with pytest.raises(ParseError, match='不存在'):
            parse_book('nonexistent.txt')

    def test_relative_path_falls_back_to_repo_root(self, tmp_path: Path, monkeypatch) -> None:
        book = tmp_path / 'fallback_book.txt'
        book.write_text('第一章 开端\n正文内容。', encoding='utf-8')
        monkeypatch.chdir(tmp_path.parent)
        monkeypatch.setattr(
            'reading_assistant.parsers.base.get_abs_path', lambda rel: str(tmp_path / rel)
        )

        result = parse_book('fallback_book.txt')

        assert result.title == 'fallback_book'
        assert result.chapter_count == 1

    def test_parse_repo_sample_from_other_cwd(self, tmp_path: Path, monkeypatch) -> None:
        monkeypatch.chdir(tmp_path)

        result = parse_book('data/books/sample_book.txt')

        assert result.chapter_count == 3
